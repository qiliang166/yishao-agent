# -*- coding: utf-8 -*-
"""R2-H 验收：docx/xlsx → Markdown 导入端点"""
import json, io, sys, uuid
import urllib.request

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8", errors="replace")
BASE = "http://localhost:8766"
TOKEN = open("dev-log/bk_admin_token").read().strip()

results = []
def check(name, ok, extra=""):
    results.append(ok)
    print(("PASS" if ok else "FAIL"), name, extra)

def upload(filename, data):
    boundary = uuid.uuid4().hex
    body = io.BytesIO()
    body.write(f"--{boundary}\r\n".encode())
    body.write(f'Content-Disposition: form-data; name="file"; filename="{filename}"\r\n'.encode())
    body.write(b"Content-Type: application/octet-stream\r\n\r\n")
    body.write(data)
    body.write(f"\r\n--{boundary}--\r\n".encode())
    r = urllib.request.Request(BASE + "/api/booklets/import-file", data=body.getvalue(), method="POST",
        headers={"Authorization": "Bearer " + TOKEN,
                 "Content-Type": f"multipart/form-data; boundary={boundary}"})
    try:
        with urllib.request.urlopen(r) as resp:
            return resp.status, json.loads(resp.read().decode("utf-8"))
    except urllib.error.HTTPError as e:
        try:
            return e.code, json.loads(e.read().decode("utf-8"))
        except Exception:
            return e.code, {}

# ── 生成真实 docx：标题/加粗/列表/表格/图片 ──
import docx
from docx.shared import Inches
from PIL import Image

doc = docx.Document()
doc.add_heading("食品安全导入测试", level=1)
p = doc.add_paragraph()
p.add_run("这是普通文字，")
p.add_run("这是加粗重点").bold = True
p.add_run("，后接普通文字。")
doc.add_heading("二级标题", level=2)
doc.add_paragraph("列表项一", style="List Bullet")
doc.add_paragraph("列表项二", style="List Bullet")
t = doc.add_table(rows=2, cols=3)
for j, v in enumerate(["岗位", "职责", "频次"]):
    t.rows[0].cells[j].text = v
for j, v in enumerate(["食品安全员", "日管控检查", "每日"]):
    t.rows[1].cells[j].text = v
img = Image.new("RGB", (80, 40), (200, 30, 30))
img_buf = io.BytesIO(); img.save(img_buf, "PNG"); img_buf.seek(0)
doc.add_picture(img_buf, width=Inches(1))
docx_buf = io.BytesIO(); doc.save(docx_buf)

st, r = upload("测试文档.docx", docx_buf.getvalue())
md = r.get("markdown", "")
check("docx 导入 200", st == 200, f"status={st} len={len(md)}")
check("docx 标题转 #", "# 食品安全导入测试" in md and "## 二级标题" in md)
check("docx 加粗转 **", "**这是加粗重点**" in md)
check("docx 列表转 -", "- 列表项一" in md and "- 列表项二" in md)
check("docx 表格转 md", "| 岗位 | 职责 | 频次 |" in md and "| 食品安全员 |" in md)
check("docx 图片转 base64", "![图片](data:image/png;base64," in md)

# ── 生成真实 xlsx：两个 sheet ──
import openpyxl
wb = openpyxl.Workbook()
ws1 = wb.active; ws1.title = "供应商"
ws1.append(["名称", "资质", "评级"])
ws1.append(["供应商A", "SC12345", "优"])
ws1.append(["供应商B", "SC67890", "良"])
ws2 = wb.create_sheet("温度记录")
ws2.append(["日期", "冷藏(℃)", "冷冻(℃)"])
ws2.append(["07-17", 4, -18])
xlsx_buf = io.BytesIO(); wb.save(xlsx_buf)

st, r = upload("测试台账.xlsx", xlsx_buf.getvalue())
md2 = r.get("markdown", "")
check("xlsx 导入 200", st == 200, f"status={st} len={len(md2)}")
check("xlsx 多sheet分节", "### 供应商" in md2 and "### 温度记录" in md2)
check("xlsx 表格内容", "| 名称 | 资质 | 评级 |" in md2 and "| 供应商A | SC12345 | 优 |" in md2)
check("xlsx 数字单元格", "| 07-17 | 4 | -18 |" in md2)

# ── 边界 ──
st, r = upload("老文档.doc", b"\xd0\xcf\x11\xe0old-doc-binary")
check(".doc 友好提示 400", st == 400 and "另存为" in r.get("detail", ""), f"status={st}")
st, r = upload("evil.exe", b"MZ....")
check("非白名单类型 400", st == 400, f"status={st}")
st, r = upload("空.docx", b"")
check("空文件 400", st == 400, f"status={st}")
st, r = upload("坏.docx", b"not-a-zip-at-all")
check("坏 docx 400", st == 400 and "无法解析" in r.get("detail", ""), f"status={st}")

print(f"\n== {sum(results)}/{len(results)} PASS ==")
